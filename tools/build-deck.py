#!/usr/bin/env python3
"""Build the ContribMatch pitch deck (docs/ContribMatch-Pitch-Deck.pptx).

The deck follows the 13-slide outline in docs/DEMO.md and uses the project's
own design tokens (paper / ink / one accent). Screenshots come from
docs/screenshots. Run:  python tools/build-deck.py  (needs python-pptx).

The generated file is committed, so editing the deck does not require Python:
regenerate only when the outline or the screenshots change.
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Pt

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "ContribMatch-Pitch-Deck.pptx"

PAPER = RGBColor(0xF2, 0xF1, 0xEE)
PAPER_RAISED = RGBColor(0xFA, 0xF9, 0xF7)
PAPER_SUNKEN = RGBColor(0xE9, 0xE8, 0xE4)
INK = RGBColor(0x10, 0x10, 0x14)
INK_2 = RGBColor(0x66, 0x66, 0x67)
ACCENT = RGBColor(0x2D, 0x81, 0xFF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SLIDE_W = 13.333
SLIDE_H = 7.5

prs = Presentation()
prs.slide_width = int(SLIDE_W * 914400)
prs.slide_height = int(SLIDE_H * 914400)
blank = prs.slide_layouts[6]


def add_slide():
    slide = prs.slides.add_slide(blank)
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = PAPER
    return slide


def rect(slide, x, y, w, h, color, line=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, int(x * 914400), int(y * 914400), int(w * 914400), int(h * 914400))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    if not line:
        shape.line.fill.background()
    return shape


def textbox(slide, x, y, w, h, lines, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT, font="Calibri", space_after=6):
    box = slide.shapes.add_textbox(int(x * 914400), int(y * 914400), int(w * 914400), int(h * 914400))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(lines):
        if isinstance(item, tuple):
            text, opts = item
        else:
            text, opts = item, {}
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = opts.get("align", align)
        p.space_after = Pt(opts.get("space", space_after))
        run = p.add_run()
        run.text = text
        run.font.size = Pt(opts.get("size", size))
        run.font.bold = opts.get("bold", bold)
        run.font.color.rgb = opts.get("color", color)
        run.font.name = opts.get("font", font)
    return box


def chrome(slide, label, number):
    """The slide furniture: a kicker label and a page number."""
    rect(slide, 0, 0, SLIDE_W, 0.06, ACCENT)
    textbox(slide, 0.55, 6.95, 8.0, 0.4, [(label, {"size": 10, "color": INK_2})], font="Consolas")
    textbox(slide, 12.1, 6.95, 0.8, 0.4, [(str(number).zfill(2), {"size": 10, "color": INK_2})], font="Consolas", align=PP_ALIGN.RIGHT)


def title_and_body(slide, title_sublines, bullets, kicker, number, font_size=18):
    textbox(slide, 0.55, 0.45, 12.2, 1.4, [(title_sublines, {"size": 30, "bold": True})])
    textbox(slide, 0.55, 1.9, 6.4, 4.8, bullets, size=font_size)
    chrome(slide, kicker, number)


def bullets(items):
    return [(f"•  {t}", {}) if isinstance(t, str) else t for t in items]


# The page captures are full-height scrolls; a slide only wants the
# above-the-fold part. The bottom is cropped with picture crop attributes,
# so the builder needs no derived image files.
SCREENSHOT_H = {
    "index-1280-light.png": 5104,
    "app-1280-light.png": 1593,
    "issue-1280-light.png": 1110,
    "design-system-1280-light.png": 6005,
    "open-innovation-1280-light.png": 2561,
    "packs-1280-light.png": 2906,
    "shortlist-1280-light.png": 1181,
}


def picture_top(slide, name, x, y, w, vis_h):
    """Embed the top `vis_h` pixels of a tall page screenshot."""
    pic = slide.shapes.add_picture(
        str(ROOT / "docs" / "screenshots" / name),
        int(x * 914400), int(y * 914400), width=int(w * 914400))
    pic.crop_bottom = max(0.0, 1.0 - (vis_h / SCREENSHOT_H[name]))
    pic.height = int(w * (vis_h / 1280.0) * 914400)
    return pic


# ---------------------------------------------------------------- 1 · Title
s = add_slide()
rect(s, 0, 0, SLIDE_W, SLIDE_H, INK)
rect(s, 0.7, 2.05, 0.55, 0.07, ACCENT)
textbox(s, 0.7, 2.3, 11.9, 1.9, [
    ("ContribMatch", {"size": 60, "bold": True, "color": WHITE}),
    ("Find an issue you can actually finish.", {"size": 26, "color": PAPER}),
])
textbox(s, 0.7, 4.6, 11.9, 1.6, [
    ("A name, an issue, a first pull request.", {"size": 16, "color": RGBColor(0xAA, 0xAA, 0xB0)}),
    ("Powered by open-weight Gemma · MIT licensed · Built for Hacktoberfest HackDays", {"size": 14, "color": RGBColor(0x8A, 0x8A, 0x90)}),
    ("Built by Tanmay Tiwari", {"size": 14, "color": RGBColor(0x8A, 0x8A, 0x90)}),
])
chrome(s, "contribmatch · pitch deck", 1)

# ------------------------------------------------------ 2 · The bottleneck
s = add_slide()
title_and_body(s, "Open source does not have an issue shortage.\nIt has a first-step shortage.", [
    "Is this issue still open?",
    "Is it still unclaimed?",
    "Is the project still alive?",
    "Is it a first contribution for me, in a language I use?",
    "What is the first pull request supposed to look like?",
], "the bottleneck", 2, 20)
textbox(s, 0.55, 5.9, 12.0, 0.6, [("None of these are code questions — and all of them stop people.", {"size": 15, "color": ACCENT, "bold": True})])
chrome(s, "the bottleneck", 2)

# -------------------------------------------------------- 3 · The product
s = add_slide()
title_and_body(s, "One screen. Four answers.", [], "the product", 3)
cards = [
    ("Match", "the issues a beginner can actually finish"),
    ("Score", "0–100, every point traceable to a rule"),
    ("Plan", "the files, the steps, the checklist"),
    ("Guard", "the quality checks a maintainer would run"),
]
positions = [(0.55, 1.75), (3.75, 1.75), (0.55, 3.35), (3.75, 3.35)]
for (head, body), (x, y) in zip(cards, positions):
    rect(s, x, y, 3.05, 1.45, PAPER_RAISED)
    rect(s, x, y, 3.05, 0.07, ACCENT)
    textbox(s, x + 0.25, y + 0.15, 2.55, 0.5, [(head, {"size": 20, "bold": True})])
    textbox(s, x + 0.25, y + 0.62, 2.55, 0.75, [(body, {"size": 12, "color": INK_2})])
picture_top(s, "index-1280-light.png", 7.0, 1.75, 5.8, 880)
textbox(s, 0.55, 5.05, 6.3, 1.2, [
    ("The home page: the 3D contribution wall and the activity panel, in one view.", {"size": 13, "color": INK_2}),
    ("Type a GitHub username and the matcher runs the whole pipeline live.", {"size": 13, "color": INK_2}),
])
chrome(s, "the product", 3)

# ------------------------------------------ 4 · Why open innovation matters
s = add_slide()
title_and_body(s, "Why open innovation matters", [], "why open innovation", 4)
textbox(s, 0.55, 1.9, 6.4, 4.8, bullets([
    "The contribution pipeline is the product, and the product should be owned by the people building it.",
    "The friction is not technical. It is informational. Information is something a tool can carry.",
    "No backend, no account, no data collection: the trust argument is the architecture, not a promise.",
    "MIT licensed and forkable in an afternoon, so a community can run their own version.",
]), size=19)
picture_top(s, "open-innovation-1280-light.png", 7.15, 1.9, 5.6, 1000)
chrome(s, "why open innovation matters", 4)

# ------------------------------------------------------- 5 · Architecture
s = add_slide()
title_and_body(s, "Architecture", [], "architecture", 5)
textbox(s, 0.55, 1.9, 6.4, 4.8, bullets([
    "Static site, vanilla ES modules, no build step, no framework",
    "Business logic is DOM-free; only js/ui touches the document",
    "One provider interface, two implementations, one mock",
    "Rules first: the model only ever sees the top eight",
]), size=19)
flow = [
    ("Profile Agent", "public activity → skills, level"),
    ("Scout Agent", "GitHub search → candidate issues"),
    ("Ranker", "rules → top 8 → Gemma → 0–100 + reasons"),
    ("Coach Agent", "issue → plan, comment, checklist"),
    ("Quality Guard", "the flags"),
]
y = 1.95
for name, what in flow:
    rect(s, 7.15, y, 5.6, 0.78, PAPER_RAISED)
    textbox(s, 7.4, y + 0.09, 5.2, 0.35, [(name, {"size": 14, "bold": True})])
    textbox(s, 7.4, y + 0.42, 5.2, 0.3, [(what, {"size": 11, "color": INK_2})])
    y += 0.92
rect(s, 8.1, y - 0.14, 3.7, 0.04, ACCENT)
textbox(s, 7.4, y + 0.0, 5.4, 0.3, [("shortlisting → localStorage (exports as Markdown)", {"size": 11, "color": ACCENT})])
chrome(s, "architecture", 5)

# ---------------------------------------------------------- 6 · The agents
s = add_slide()
title_and_body(s, "The agents", [], "the agents", 6)
rows = [
    ("Profile", "public repos, languages, topics, activity", "skills, level"),
    ("Scout", "GitHub search", "a candidate list"),
    ("Ranker", "rules, then the model on the top 8", "0–100 and the reasons"),
    ("Coach", "one issue", "the plan, the comment, the checklist"),
    ("Quality Guard", "the plan and the draft", "the flags"),
]
cols = [0.55, 3.3, 8.0]
heads = ["Agent", "Reads", "Writes"]
for cx, h in zip(cols, heads):
    textbox(s, cx, 2.0, 4.4, 0.4, [(h.upper(), {"size": 11, "color": INK_2, "font": "Consolas"})])
rect(s, 0.55, 2.45, 12.2, 0.02, RGBColor(0xD8, 0xD6, 0xD2))
y = 2.65
for a, r, w in rows:
    textbox(s, cols[0], y, 2.6, 0.7, [(a, {"size": 16, "bold": True})])
    textbox(s, cols[1], y, 4.5, 0.7, [(r, {"size": 15, "color": INK})])
    textbox(s, cols[2], y, 4.7, 0.7, [(w, {"size": 15, "color": ACCENT})])
    y += 0.82
textbox(s, 0.55, y + 0.15, 12.2, 0.5, [("Every agent is async run(input, ctx) and returns schema-validated JSON.", {"size": 14, "color": INK_2})])
chrome(s, "the agents", 6)

# ------------------------------------------------------ 7 · Open-weight Gemma
s = add_slide()
title_and_body(s, "Open-weight Gemma", [], "open-weight gemma", 7)
textbox(s, 0.55, 1.9, 6.4, 4.8, bullets([
    "Ollama runs Gemma locally — the model that is on your machine",
    "Google AI Studio hosts Gemma, for a shared key",
    "OLLAMA_ORIGINS is the setting people get wrong; it is in the docs",
    "The model is never the only path. Rules carry the ranking; the model carries the wording.",
    "The model rating is blended at 35% into the final score — never 100%",
    "A key is never committed and never sent anywhere but the API it belongs to",
]), size=17)
rect(s, 7.15, 1.9, 5.6, 4.4, PAPER_RAISED)
textbox(s, 7.45, 2.15, 5.0, 0.5, [("RANKING SOURCES", {"size": 11, "color": INK_2, "font": "Consolas"})])
textbox(s, 7.45, 2.7, 5.0, 3.4, [
    ("65%", {"size": 40, "bold": True, "color": INK}),
    ("transparent rules: language match, repo health, labels, freshness, activity, body hygiene", {"size": 13, "color": INK_2, "space": 12}),
    ("35%", {"size": 40, "bold": True, "color": ACCENT}),
    ("Gemma's rating of the top 8: ease, clarity, fit", {"size": 13, "color": INK_2}),
])
chrome(s, "open-weight gemma", 7)

# ----------------------------------------- 8 · Exact vs approximate data
s = add_slide()
title_and_body(s, "Exact versus approximate data", [], "honesty", 8)
modes = [
    ("Exact", "a token is set", "the full calendar from GraphQL"),
    ("Approximate", "no token", "about 90 days of public events, mapped"),
    ("Sample", "no name yet", "clearly labelled sample data"),
]
y = 2.0
for label, when, what in modes:
    rect(s, 0.55, y, 12.2, 1.05, PAPER_RAISED)
    textbox(s, 0.9, y + 0.28, 2.6, 0.5, [(label, {"size": 17, "bold": True})])
    textbox(s, 3.6, y + 0.31, 3.4, 0.5, [(when, {"size": 14, "color": INK_2})])
    textbox(s, 7.1, y + 0.31, 5.4, 0.5, [(what, {"size": 14})])
    y += 1.2
textbox(s, 0.55, y + 0.15, 12.2, 0.8, [
    ("The 3D wall always covers the same window as the panel. An approximation is never drawn as a full year.", {"size": 15, "bold": True, "color": ACCENT}),
])
chrome(s, "honesty as a feature", 8)

# ------------------------------------------------------- 9 · The interface
s = add_slide()
title_and_body(s, "The interface", [], "the interface", 9)
textbox(s, 0.55, 1.9, 5.6, 4.6, bullets([
    "Neutral, square, hairlines, one accent",
    "Depth from perspective and flat offset shadows, never from glow",
    "English, Hindi and Hinglish — including the generated notes",
    "Every control has six states and a 24px minimum target",
    "WCAG AA throughout; the audit runs on every page, in both themes",
]), size=16)
picture_top(s, "design-system-1280-light.png", 6.35, 1.9, 6.4, 1150)
chrome(s, "the interface", 9)

# ------------------------------------------------- 10 · Community packs
s = add_slide()
title_and_body(s, "Community packs", [], "community packs", 10)
rect(s, 0.55, 1.9, 6.4, 1.0, PAPER_SUNKEN)
textbox(s, 0.8, 2.05, 6.0, 0.7, [('{ "id": "web-basics", "type": "skill",', {"size": 13, "color": INK, "font": "Consolas", "space": 0}), ('  "skills": ["html", "css"] }', {"size": 13, "color": INK, "font": "Consolas", "space": 0})])
textbox(s, 0.55, 3.2, 6.4, 3.2, bullets([
    "A pack is a plain JSON or YAML file describing skills, tone and vocabulary for an ecosystem",
    "Five ship with the project",
    "npm run validate:packs checks them",
    "Adding one is a pull request, not a ticket — this is the collaboration model, not a feature list",
]), size=16)
picture_top(s, "packs-1280-light.png", 7.15, 1.9, 5.6, 1150)
chrome(s, "community packs", 10)

# ------------------------------------------------------------- 11 · Try it
s = add_slide()
title_and_body(s, "Try it", [], "try it", 11)
rect(s, 0.55, 1.9, 6.4, 1.35, INK)
textbox(s, 0.8, 2.05, 6.0, 1.1, [
    ("git clone https://github.com/Tanmay-Tiwaricyber/contribmatch", {"size": 12, "color": PAPER, "font": "Consolas", "space": 2}),
    ("cd contribmatch && npm install && npm run serve", {"size": 12, "color": PAPER, "font": "Consolas", "space": 2}),
])
textbox(s, 0.55, 3.55, 6.4, 3.2, bullets([
    "npm test runs everything: the packs, the rules, the providers, the orchestrator, and the copy rules",
    "No key needed. Mock mode runs the whole flow offline.",
    "Gemma locally: ollama pull gemma3:4b",
]), size=16)
rect(s, 7.15, 1.9, 5.6, 4.4, PAPER_RAISED)
textbox(s, 7.45, 2.15, 5.0, 0.5, [("LINKS", {"size": 11, "color": INK_2, "font": "Consolas"})])
textbox(s, 7.45, 2.7, 5.0, 3.4, [
    ("Live app", {"size": 14, "bold": True}),
    ("https://tanmay-tiwaricyber.github.io/contribmatch/", {"size": 13, "color": ACCENT, "space": 12}),
    ("Repository", {"size": 14, "bold": True}),
    ("https://github.com/Tanmay-Tiwaricyber/contribmatch", {"size": 13, "color": ACCENT, "space": 12}),
    ("Demo video", {"size": 14, "bold": True}),
    ("link to be added on submission", {"size": 13, "color": INK_2}),
])
chrome(s, "try it", 11)

# ------------------------------------------------ 12 · By design, by choice
s = add_slide()
textbox(s, 0.55, 0.45, 12.2, 1.4, [("By design, not by accident", {"size": 30, "bold": True})])
textbox(s, 0.55, 1.05, 11.5, 0.6, [("Five decisions that shape the product. Ask us about any of them.", {"size": 15, "color": INK_2})])
rows = [
    ("The ranking is a rubric — every score traces to rules you can read, not a black-box guess."),
    ("Gemma refines and writes; the rules decide. Turn the model off and you still get a shortlist and a plan."),
    ("The graph always says what it covers — a token unlocks the exact year; without one you get 90 days of public events, clearly labelled."),
    ("Results are bounded by GitHub's own search index. Nothing is invented, and the static catalogue keeps good data flowing."),
    ("No backend and no write path: nothing is ever posted on your behalf. Every action stays yours."),
]
y = 2.05
for i, text in enumerate(rows, 1):
    textbox(s, 0.55, y + 0.02, 0.9, 0.6, [(str(i).zfill(2), {"size": 26, "bold": True, "color": INK_2, "font": "Consolas"})])
    textbox(s, 1.75, y + 0.08, 10.9, 0.7, [(text, {"size": 15})])
    if i < len(rows):
        rect(s, 1.75, y + 0.82, 10.9, 0.015, RGBColor(0xD8, 0xD6, 0xD2))
    y += 0.94
chrome(s, "design decisions", 12)

# ------------------------------------------------------------- 13 · Close
s = add_slide()
rect(s, 0, 0, SLIDE_W, SLIDE_H, INK)
rect(s, 0.7, 2.5, 0.55, 0.07, ACCENT)
textbox(s, 0.7, 2.7, 11.9, 1.2, [
    ("Find an issue you can actually finish.", {"size": 40, "bold": True, "color": WHITE}),
])
textbox(s, 0.7, 4.1, 11.9, 2.2, [
    ("Live: https://tanmay-tiwaricyber.github.io/contribmatch/   ·   Code: https://github.com/Tanmay-Tiwaricyber/contribmatch   ·   Demo: [video link]", {"size": 14, "color": PAPER, "font": "Consolas"}),
    ("MIT licensed. Open-source because the contribution pipeline is the product, and the product should belong to the people building it.", {"size": 14, "color": RGBColor(0xAA, 0xAA, 0xB0)}),
    ("Built by Tanmay Tiwari · Built for Hacktoberfest HackDays · Powered by open-weight Gemma", {"size": 14, "color": RGBColor(0x8A, 0x8A, 0x90)}),
])
chrome(s, "contribmatch · thank you", 13)

prs.save(OUT)
print(f"Saved {OUT} with {len(prs.slides)} slides")
